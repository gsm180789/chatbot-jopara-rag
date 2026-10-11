"""Carga de documentos del corpus (.docx y .pdf) a texto plano + metadatos."""
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document as DocxDocument
from docx.table import Table
from docx.text.paragraph import Paragraph
#from pypdf import PdfReader
import pdfplumber


@dataclass
class LoadedDocument:
    """Resultado de leer un archivo: texto completo y datos de origen."""
    text: str
    source: str                    # nombre del archivo
    doc_type: str                  # "docx" o "pdf"
    metadata: dict = field(default_factory=dict)


def load_docx(path: Path) -> LoadedDocument:
    """Lee un .docx respetando el orden de párrafos y tablas."""
    doc = DocxDocument(path)
    blocks: list[str] = []

    # body.iterchildren() devuelve párrafos y tablas en el orden real
    for child in doc.element.body.iterchildren():
        if child.tag.endswith("}p"):
            text = Paragraph(child, doc).text.strip()
            if text:
                blocks.append(text)
        elif child.tag.endswith("}tbl"):
            table = Table(child, doc)
            for row in table.rows:
                # celdas combinadas repiten texto: se eliminan duplicados contiguos
                cells: list[str] = []
                for cell in row.cells:
                    t = cell.text.strip().replace("\n", " ")
                    if t and (not cells or cells[-1] != t):
                        cells.append(t)
                if cells:
                    blocks.append(" | ".join(cells))

    return LoadedDocument(
        text="\n".join(blocks),
        source=path.name,
        doc_type="docx",
        metadata={"n_blocks": len(blocks), "path": str(path)},
    )


def load_pdf(path: Path) -> LoadedDocument:
    """Lee un PDF con texto seleccionable, respetando filas de tablas.

    pdfplumber agrupa las palabras por posición vertical, así cada fila
    de una tabla queda en una sola línea (pypdf las separaba por columna).
    """
    pages: list[str] = []
    with pdfplumber.open(path) as pdf:
        n_pages = len(pdf.pages)
        for page in pdf.pages:
            text = (page.extract_text(x_tolerance=2, y_tolerance=3) or "").strip()
            if text:
                pages.append(text)

    return LoadedDocument(
        text="\n\n".join(pages),
        source=path.name,
        doc_type="pdf",
        metadata={"n_pages": n_pages, "n_pages_with_text": len(pages), "path": str(path)},
    )


LOADERS = {".docx": load_docx, ".pdf": load_pdf}


def load_document(path: str | Path) -> LoadedDocument:
    """Elige el loader según la extensión del archivo."""
    path = Path(path)
    loader = LOADERS.get(path.suffix.lower())
    if loader is None:
        raise ValueError(f"Formato no soportado: {path.suffix} ({path.name})")
    return load_docx(path) if loader is load_docx else loader(path)